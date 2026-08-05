import json
from io import BytesIO
from datetime import date, datetime
from decimal import Decimal

from openpyxl import Workbook
from sqlalchemy import create_engine, select
from sqlalchemy.orm import Session

from app.database import Base
from app.models import (
    BonusRule,
    FinancialEntry,
    PayableDocument,
    PayableMovement,
    PortalBonusEvent,
    PortalBonusEventSource,
    PortalBonusMatch,
    PortalStatementImport,
    Purchase,
    PurchaseItem,
)
from app.services.ipiranga_portal import (
    ALGORITHM_VERSION,
    CREDIT_REPORT_CATEGORY,
    ParsedPortalEvent,
    POSTPAID_CATEGORY,
    _parse_sheet,
    _parse_portal_pdf_text,
    build_ipiranga_cycle_ledger,
    import_ipiranga_statement,
    match_ipiranga_events,
    parse_ipiranga_statement,
)
from app.services.reconciliation import (
    _portal_invoice_credit_evidence,
    _portal_invoice_credit_issued_evidence,
)
from app.services.reconciliation_workspace import _automatic_policy
from app.services.rules import add_months, money
from app.services.seed import seed_reference_data
from app.scripts.apply_unit_004_portal_usage_evidence import apply_unit_004_portal_usage_evidence


class FakeCell:
    def __init__(self, value):
        self.value = value
        self.ctype = 1


class FakeSheet:
    def __init__(self, rows):
        self.rows = rows
        self.nrows = len(rows)
        self.ncols = max(len(row) for row in rows)

    def cell_value(self, row, column):
        return self.rows[row][column] if column < len(self.rows[row]) else ""

    def cell(self, row, column):
        return FakeCell(self.cell_value(row, column))


def test_parser_reads_only_transaction_row_and_preserves_metadata():
    sheet = FakeSheet(
        [
            ["Extrato Consolidado"],
            [],
            ["CNPJ do cliente", 905896980001],
            ["Categoria", "Bonificação Postecipada"],
            ["Período", "01/03/25", "até", "31/08/25"],
            [],
            ["Data", "Produto", "Descrição", "Referência", "Valor (R$)", "Totalizadores (R$)", "Saldo dia (R$)"],
            ["01/05/2025", "Bonificação Postecipada", "17", "17", "17", "14.007,00", "17"],
            ["01/05/2025", "Bonificação Postecipada", "Bonificação", "17", "14.007,00", "17", "17"],
            ["01/05/2025", "Saldo da movimentação do dia", "17", "17", "17", "17", "14.007,00"],
        ]
    )
    parsed = _parse_sheet(sheet)
    assert parsed.client_cnpj == "00905896980001"
    assert parsed.period_start == date(2025, 3, 1)
    assert parsed.period_end == date(2025, 8, 31)
    assert len(parsed.events) == 1
    assert parsed.events[0].portal_date == date(2025, 5, 1)
    assert parsed.events[0].value == Decimal("14007.00")


def test_pdf_parser_reads_only_explicit_postpaid_credits():
    parsed = _parse_portal_pdf_text(
        [
            "\n".join(
                [
                    "EXTRATO CONSOLIDADO",
                    "CNPJ do cliente: 905896980004",
                    "Categoria: Todos",
                    "Per?odo: 01/08/25 at? 01/02/26",
                    "17/10/2025 Bonifica??o Postecipada 8.400,00",
                    "17/10/2025 Nota Fiscal Produto -53.619,00",
                    "18/11/2025 Bonifica??o Postecipada 8.750,00",
                    "18/11/2025 Saldo da movimenta??o do dia 8.750,00",
                ]
            )
        ]
    )
    assert parsed.client_cnpj == "00905896980004"
    assert parsed.period_start == date(2025, 8, 1)
    assert parsed.period_end == date(2026, 2, 1)
    assert [(event.portal_date, event.value) for event in parsed.events] == [
        (date(2025, 10, 17), Decimal("8400.00")),
        (date(2025, 11, 18), Decimal("8750.00")),
    ]


def test_pdf_parser_reads_texaco_own_note_credit_from_consolidated_statement():
    parsed = _parse_portal_pdf_text(
        [
            "\n".join(
                [
                    "EXTRATO CONSOLIDADO",
                    "CNPJ do cliente: 905896980005",
                    "Categoria: Todos",
                    "Periodo: 01/06/26 ate 23/07/26",
                    "16/06/2026 Nota Fiscal Produto -53.619,00",
                    "16/06/2026 Nota Pr?pria 1.080,00",
                    "16/06/2026 Saldo da movimentacao do dia 1.080,00",
                ]
            )
        ]
    )
    assert parsed.client_cnpj == "00905896980005"
    assert [(event.portal_date, event.value, event.description) for event in parsed.events] == [
        (date(2026, 6, 16), Decimal("1080.00"), "Nota Propria"),
    ]
    assert parsed.events[0].raw["portal_product_total"] == "53619.00"


def test_pdf_parser_keeps_texaco_commercial_action_separate_from_postpaid_credit():
    parsed = _parse_portal_pdf_text(
        [
            "\n".join(
                [
                    "EXTRATO CONSOLIDADO",
                    "CNPJ do cliente: 125642760001",
                    "Categoria: Todos",
                    "Periodo: 01/06/26 ate 24/07/26",
                    "23/06/2026 Acao comercial Combs 29.360,00",
                    "23/06/2026 Nota Pr?pria 920,00",
                ]
            )
        ]
    )
    assert [(event.description, event.category, event.value) for event in parsed.events] == [
        ("Acao comercial combustiveis", "supplemental_other_credit", Decimal("29360.00")),
        ("Nota Propria", POSTPAID_CATEGORY, Decimal("920.00")),
    ]


def test_import_is_idempotent_and_deduplicates_overlapping_files(monkeypatch):
    engine = create_engine("sqlite:///:memory:")
    Base.metadata.create_all(engine)
    event = ParsedPortalEvent(9, date(2025, 8, 1), Decimal("12963"), "Postecipada", "Bonificação", "17", [])
    parsed = type("Parsed", (), {
        "client_cnpj": "00905896980001",
        "category_label": POSTPAID_CATEGORY,
        "period_start": date(2025, 3, 1),
        "period_end": date(2025, 8, 31),
        "events": [event],
    })()
    monkeypatch.setattr("app.services.ipiranga_portal.parse_ipiranga_statement", lambda *_: parsed)
    with Session(engine) as db:
        seed_reference_data(db)
        first, repeated = import_ipiranga_statement(
            db, unit_code="001", filename="primeiro.xls", content_type="application/vnd.ms-excel",
            content=b"first", uploaded_by="admin",
        )
        second, second_repeated = import_ipiranga_statement(
            db, unit_code="001", filename="sobreposto.xls", content_type="application/vnd.ms-excel",
            content=b"second", uploaded_by="admin",
        )
        same_file, same_repeated = import_ipiranga_statement(
            db, unit_code="001", filename="primeiro.xls", content_type="application/vnd.ms-excel",
            content=b"first", uploaded_by="admin",
        )
        assert not repeated and not second_repeated and same_repeated
        assert first.imported_count == 1 and second.imported_count == 0
        assert second.duplicate_count == 1
        assert same_file.id == first.id


def test_credit_report_groups_postpaid_and_supplemental_events():
    workbook = Workbook()
    sheet = workbook.active
    sheet.append([
        "Nº do Cliente",
        "Número Documen.",
        "Tipo Doc.",
        "Data Vncto",
        "Tipo de Documento",
        "Cia Doc",
        "Item Pgto",
        "Data Fatura",
        "Valor Bruto",
        "Valor Aberto",
        "D.Venc Desc.",
        "Dias após Vcto",
        "Nº Avisos",
        "Data Fech.",
        "Cntr G/L",
        "Md. Bs.",
        "Código PAI",
        "Taxa Câmbio",
        "Moeda Trans.",
        "Cód. Ded.",
        "Cód. Motivo Dedução",
        "C C",
        "Unidade Negócios",
        "Revisão Efetuada",
        "Data Estor.",
        "Status Pgto",
        "Cód. Status Pgto",
        "Nome-Observação",
        "Descrição do Cliente",
        "Descrição Cadastro Pai",
        "Número Pagador",
        "Descrição Cad. Geral do Pagador",
        "Valor Md. Estr.",
        "Valor em Aberto Moeda Estr.",
        "Desc. Disp. Md. Estr.",
        "Foreign Disc Obtido",
        "I P",
        "Data G/L",
        "Número Lote",
        "Tipo Lote",
        "Descrição Tipo Lote",
        "Data Lote",
        "Cia",
        "Documento Original",
    ])
    sheet.append([2062423, 194906, "CU", date(2025, 4, 27), "Abatimento Concedido", "07090", "001", date(2025, 4, 27), -4668.95, None, date(2025, 4, 27), None, None, date(2025, 5, 1), "CLI", "BRL", 2886366, None, "BRL", None, None, "D", "070346", None, None, "P", "Pagamento Total", "Bonificação", "GBI", "PAI", 2062423, "GBI", None, None, None, None, "1", date(2025, 4, 27), 43807953, "IB", "Faturas", date(2025, 4, 27), "00007", 194906])
    sheet.append([2062423, 194906, "CU", date(2025, 4, 27), "Abatimento Concedido", "07090", "002", date(2025, 4, 27), -4668.96, None, date(2025, 4, 27), None, None, date(2025, 5, 1), "CLI", "BRL", 2886366, None, "BRL", None, None, "D", "070346", None, None, "P", "Pagamento Total", "Bonificação", "GBI", "PAI", 2062423, "GBI", None, None, None, None, "1", date(2025, 4, 27), 43807953, "IB", "Faturas", date(2025, 4, 27), "00007", 194906])
    sheet.append([2062423, 194906, "CU", date(2025, 4, 27), "Abatimento Concedido", "07090", "003", date(2025, 4, 27), -4669.09, None, date(2025, 4, 27), None, None, date(2025, 5, 1), "CLI", "BRL", 2886366, None, "BRL", None, None, "D", "070346", None, None, "P", "Pagamento Total", "Bonificação", "GBI", "PAI", 2062423, "GBI", None, None, None, None, "1", date(2025, 4, 27), 43807953, "IB", "Faturas", date(2025, 4, 27), "00007", 194906])
    sheet.append([2062423, 272405, "CZ", date(2024, 12, 26), "Aviso de Crédito depósito aut.", "07090", "001", date(2024, 12, 26), -5626.42, None, date(2024, 12, 26), None, None, date(2024, 12, 27), "CLI", "BRL", 2886366, None, "BRL", None, None, "D", "070346", None, None, "P", "Pagamento Total", "PAGTO - 26/12/2024", "GBI", "PAI", 2062423, "GBI", None, None, None, None, "1", date(2024, 12, 26), 43807953, "IB", "Faturas", date(2024, 12, 27), "00007", 272405])
    sheet.append([2062423, 316554, "CR", date(2025, 2, 17), "Aviso de Credito", "07090", "001", date(2025, 2, 17), -27348.50, None, date(2025, 2, 17), None, None, date(2025, 3, 11), "CLI", "BRL", 2886366, None, "BRL", None, None, "D", "071000", None, None, "P", "Pagamento Total", "Invoice341WNS287545517022025", "GBI", "PAI", 2062423, "GBI", None, None, None, None, "1", date(2025, 2, 17), 39183945, "IB", "Faturas", date(2025, 3, 11), "00007", 316554])

    buffer = BytesIO()
    workbook.save(buffer)
    parsed = parse_ipiranga_statement(buffer.getvalue(), "relatorio_creditos.xlsx")

    assert parsed.category_label == CREDIT_REPORT_CATEGORY
    assert parsed.period_start == date(2024, 12, 26)
    assert parsed.period_end == date(2025, 4, 27)
    assert len(parsed.events) == 3

    by_category = {event.category: event for event in parsed.events}
    postpaid = by_category[POSTPAID_CATEGORY]
    assert postpaid.portal_date == date(2025, 5, 1)
    assert postpaid.value == Decimal("14007.00")
    assert postpaid.row_label == "linhas 2-4"
    assert postpaid.raw["row_numbers"] == [2, 3, 4]

    auto_deposit = by_category["supplemental_auto_deposit"]
    assert auto_deposit.portal_date == date(2024, 12, 27)
    assert auto_deposit.value == Decimal("5626.42")

    credit_notice = by_category["supplemental_credit_notice"]
    assert credit_notice.portal_date == date(2025, 3, 11)
    assert credit_notice.value == Decimal("27348.50")


def test_contract_parcels_report_reads_only_issued_credits():
    workbook = Workbook()
    sheet = workbook.active
    headers = [
        "Razão Social", "CNPJ", "Parcela", "Status", "Volume Real", "Volume a Atingir",
        "Performance Real %", "Período Inicial", "Período Final", "Data da Emissão",
        "Valor Parcela", "Valor Calculado Parcela", "Índice Correcao", "Valor Crédito",
    ]
    sheet.append(headers)
    sheet.append([
        "Gbi Combustiveis Ltda.", "905896980008", "01", "Crédito Emitido", "110", "100", "110%",
        "27/06/2022", "26/07/2022", "01/08/2022", "R$ 5.000,00", "R$ 5.500,00", "", "R$ 5.500,00",
    ])
    sheet.append([
        "Gbi Combustiveis Ltda.", "905896980008", "02", "Em Andamento", "55", "100", "55%",
        "27/07/2022", "26/08/2022", "", "R$ 5.000,00", "R$ 2.750,00", "", "-",
    ])
    buffer = BytesIO()
    workbook.save(buffer)

    parsed = parse_ipiranga_statement(buffer.getvalue(), "parcelas-008.xlsx")

    assert parsed.client_cnpj == "00905896980008"
    assert parsed.period_start == date(2022, 6, 27)
    # Scheduled installments must not make an imported statement look newer
    # than its last issued credit.
    assert parsed.period_end == date(2022, 7, 26)
    assert len(parsed.events) == 1
    event = parsed.events[0]
    assert event.portal_date == date(2022, 8, 1)
    assert event.value == Decimal("5500.00")
    assert event.reference == "Parcela 01"
    assert event.raw["business_classification"] == "postpaid_issued"


def test_008_seed_has_the_confirmed_postpaid_rule():
    engine = create_engine("sqlite:///:memory:")
    Base.metadata.create_all(engine)
    with Session(engine) as db:
        seed_reference_data(db)
        rule = db.scalar(
            select(BonusRule).where(
                BonusRule.unit_code == "008",
                BonusRule.company_code == "IPIRANGA",
                BonusRule.kind == "distributor_credit",
            )
        )
        assert rule.effective_from == date(2022, 6, 27)
        assert rule.effective_to == date(2027, 6, 26)
        assert rule.rate_per_liter == Decimal("0.05")


def test_contract_parcel_credit_is_confirmed_as_issued_not_unmatched():
    engine = create_engine("sqlite:///:memory:")
    Base.metadata.create_all(engine)
    with Session(engine) as db:
        event = PortalBonusEvent(
            id="parcel-008-1",
            unit_code="008",
            company_code="IPIRANGA",
            category=POSTPAID_CATEGORY,
            portal_date=date(2022, 8, 1),
            value=Decimal("5500"),
            description="Credito Emitido",
            business_classification="postpaid_issued",
            event_key="parcel-008-key-1",
        )
        db.add(event)
        db.flush()

        assert match_ipiranga_events(db, "008") == 1
        match = db.scalar(select(PortalBonusMatch).where(PortalBonusMatch.event_id == event.id))
        assert match.status == "issued"
        assert "Crédito emitido explicitamente" in match.match_basis


def test_ipiranga_matcher_preserves_unit_004_portal_usage_confirmation():
    """A direct portal detail must not be replaced by an ERP proximity guess."""
    engine = create_engine("sqlite:///:memory:")
    Base.metadata.create_all(engine)
    with Session(engine) as db:
        event = PortalBonusEvent(
            id="unit-004-usage",
            unit_code="004",
            company_code="IPIRANGA",
            category=POSTPAID_CATEGORY,
            portal_date=date(2026, 1, 16),
            value=Decimal("10360.01"),
            description="BonificaÃ§Ã£o Postecipada",
            event_key="unit-004-usage-key",
        )
        db.add(event)
        db.flush()
        db.add(
            PortalBonusMatch(
                event_id=event.id,
                status="portal_usage_confirmed",
                purchase_entry_id=900090260,
                match_basis="Detalhe do portal informa a NF utilizada.",
                details_json='{"usage_invoice_number":"3008303"}',
                algorithm_version=ALGORITHM_VERSION,
            )
        )
        db.commit()

        assert match_ipiranga_events(db, "004") == 1

        match = db.scalar(select(PortalBonusMatch).where(PortalBonusMatch.event_id == event.id))
        assert match.status == "portal_usage_confirmed"
        assert match.purchase_entry_id == 900090260


def test_unit_004_portal_usage_evidence_registers_the_explicit_used_invoice():
    engine = create_engine("sqlite:///:memory:")
    Base.metadata.create_all(engine)
    with Session(engine) as db:
        purchase = Purchase(
            erp_entry_id=900090260,
            unit_code="004",
            supplier_name="IPIRANGA PRODUTOS DE PETROLEO SA",
            supplier_cnpj="33337122015906",
            mapped_company_code="IPIRANGA",
            invoice_number="3008303",
            access_key="43260133337122015906550030030083031951343959",
            purchase_date=date(2026, 1, 13),
            total_liters=Decimal("20000"),
            s10_liters=Decimal("0"),
            gross_value=Decimal("108650.00"),
            net_value=Decimal("108650.00"),
        )
        event = PortalBonusEvent(
            id="usage-004-jan",
            event_key="usage-004-jan",
            unit_code="004",
            company_code="IPIRANGA",
            category="postpaid",
            portal_date=date(2026, 1, 16),
            value=Decimal("10360.01"),
        )
        db.add_all([purchase, event])
        db.flush()

        result = apply_unit_004_portal_usage_evidence(
            db,
            evidence_rows=[
                {
                    "portal_date": date(2026, 1, 16),
                    "credit_value": Decimal("10360.01"),
                    "purchase_entry_id": 900090260,
                    "invoice_number": "3008303",
                    "gross_value": Decimal("108650.00"),
                    "access_key": "43260133337122015906550030030083031951343959",
                }
            ],
        )

        match = db.scalar(select(PortalBonusMatch).where(PortalBonusMatch.event_id == event.id))
        assert result["linked"] == 1
        assert match.status == "portal_usage_confirmed"
        assert match.purchase_entry_id == purchase.erp_entry_id
        assert json.loads(match.details_json)["usage_invoice_number"] == "3008303"


def _add_exact_chain(
    db: Session,
    event_value: Decimal = Decimal("14007"),
    *,
    category: str = "postpaid",
    portal_date: date = date(2025, 5, 1),
    event_id: str = "portal-1",
    event_key: str = "event-1",
):
    purchase = Purchase(
        erp_entry_id=9001, unit_code="001", supplier_name="IPIRANGA PRODUTOS",
        supplier_cnpj="33337122015906", mapped_company_code="IPIRANGA", invoice_number="2904663",
        purchase_date=date(2025, 4, 29), invoice_issue_date=date(2025, 4, 25),
        total_liters=Decimal("5000"), s10_liters=Decimal("0"), gross_value=1, net_value=1,
    )
    db.add(purchase)
    document = PayableDocument(
        unit_code="001", person_id=44, title_type="NF", document_id="904663", sequence="01",
        erp_entry_id=9001, invoice_number="2904663", document_value=Decimal("30000"), other_discount=0,
        issue_date=date(2025, 4, 25), due_date=date(2025, 4, 30), payment_date=date(2025, 4, 30), balance=0,
    )
    db.add(document)
    db.add_all(
        [
            PayableMovement(
                erp_key="B:9001", unit_code="001", person_id=44, title_type="NF", document_id="904663",
                document_sequence="01", movement_type="B", movement_date=date(2025, 4, 30), amount=Decimal("30000"),
                payment_sequence=1, reversal_sequence=0, financial_launch_id=7001, financial_sequence=1,
            ),
            PayableMovement(
                erp_key="D:9001", unit_code="001", person_id=44, title_type="NF", document_id="904663",
                document_sequence="01", movement_type="D", movement_date=date(2025, 4, 30), amount=event_value,
                payment_sequence=1, reversal_sequence=0, financial_launch_id=7001, financial_sequence=1,
            ),
        ]
    )
    db.add(
        FinancialEntry(
            erp_launch_id=7001, erp_sequence=2, unit_code="001", entry_date=date(2025, 4, 30),
            history_code=6204, direction="E", value=event_value, document_id="904663",
            history_text="DESCONTOS OBTIDOS CFE DOCUM. IPIRANGA", origin="BPag",
        )
    )
    event = PortalBonusEvent(
        id=event_id, unit_code="001", company_code="IPIRANGA", category=category,
        portal_date=portal_date, value=event_value, description="Bonificação", event_key=event_key,
    )
    db.add(event)
    db.flush()
    return event


def test_portal_event_requires_the_complete_unique_chain():
    engine = create_engine("sqlite:///:memory:")
    Base.metadata.create_all(engine)
    with Session(engine) as db:
        event = _add_exact_chain(db)
        assert match_ipiranga_events(db) == 1
        match = db.scalar(select(PortalBonusMatch).where(PortalBonusMatch.event_id == event.id))
        assert match.status == "exact"
        assert match.movement_key == "D:9001"
        assert match.purchase_entry_id == 9001
        assert match.financial_entry_id is not None


def test_supplemental_portal_event_can_share_the_same_exact_chain_logic():
    engine = create_engine("sqlite:///:memory:")
    Base.metadata.create_all(engine)
    with Session(engine) as db:
        event = _add_exact_chain(
            db,
            category="supplemental_auto_deposit",
            portal_date=date(2025, 4, 30),
            event_id="portal-2",
            event_key="event-2",
        )
        assert match_ipiranga_events(db) == 1
        match = db.scalar(select(PortalBonusMatch).where(PortalBonusMatch.event_id == event.id))
        assert match.status == "exact"
        assert match.movement_key == "D:9001"
        assert match.purchase_entry_id == 9001


def test_texaco_portal_own_note_requires_formula_and_unique_paid_title():
    engine = create_engine("sqlite:///:memory:")
    Base.metadata.create_all(engine)
    with Session(engine) as db:
        seed_reference_data(db)
        purchase = Purchase(
            erp_entry_id=5005, unit_code="005", supplier_name="IPIRANGA PRODUTOS",
            supplier_cnpj="33337122015906", mapped_company_code="TEXACO", invoice_number="3000001",
            purchase_date=date(2026, 6, 10), total_liters=Decimal("10000"), s10_liters=0,
            gross_value=Decimal("100000"), net_value=Decimal("100000"),
        )
        purchase.items.append(PurchaseItem(
            item_code="1", sequence=1, description="GASOLINA", unit="L", quantity=Decimal("10000"),
            unit_value=Decimal("10"), total_value=Decimal("100000"),
        ))
        db.add(purchase)
        db.add(PayableDocument(
            unit_code="005", person_id=99, title_type="NF", document_id="000001", sequence="01",
            erp_entry_id=5005, invoice_number="3000001", document_value=Decimal("100000"), other_discount=0,
            issue_date=date(2026, 6, 10), due_date=date(2026, 6, 15), payment_date=date(2026, 6, 16), balance=0,
        ))
        db.add(PayableMovement(
            erp_key="B:5005", unit_code="005", person_id=99, title_type="NF", document_id="000001",
            document_sequence="01", movement_type="B", movement_date=date(2026, 6, 16), amount=Decimal("100000"),
            payment_sequence=1, reversal_sequence=0, financial_launch_id=5005, financial_sequence=1,
        ))
        db.add(FinancialEntry(
            erp_launch_id=5005, erp_sequence=1, unit_code="005", entry_date=date(2026, 6, 16),
            history_code=51, direction="S", value=Decimal("100000"), document_id="000001", origin="BPag",
        ))
        event = PortalBonusEvent(
            id="texaco-portal-1", unit_code="005", company_code="TEXACO", category="postpaid",
            portal_date=date(2026, 6, 16), value=Decimal("400"), description="Nota Propria", event_key="texaco-key-1",
        )
        db.add(event)
        db.flush()
        assert match_ipiranga_events(db, "005", company_code="TEXACO") == 1
        match = db.scalar(select(PortalBonusMatch).where(PortalBonusMatch.event_id == event.id))
        assert match.status == "portal_exact"
        assert match.purchase_entry_id == 5005


def _add_texaco_paid_purchase(
    db: Session,
    *,
    entry_id: int,
    invoice_number: str,
    document_id: str,
    liters: str,
    purchase_date: date,
    payment_date: date | None,
    due_date: date | None = None,
):
    value = Decimal(liters) * Decimal("5")
    purchase = Purchase(
        erp_entry_id=entry_id,
        unit_code="777",
        supplier_name="IPIRANGA PRODUTOS",
        supplier_cnpj="33337122015906",
        mapped_company_code="TEXACO",
        invoice_number=invoice_number,
        purchase_date=purchase_date,
        total_liters=Decimal(liters),
        s10_liters=Decimal("0"),
        gross_value=value,
        net_value=value,
    )
    purchase.items.append(PurchaseItem(
        item_code="1",
        sequence=1,
        description="GASOLINA",
        unit="L",
        quantity=Decimal(liters),
        unit_value=Decimal("5"),
        total_value=value,
    ))
    db.add(purchase)
    db.add(PayableDocument(
        unit_code="777",
        person_id=77,
        title_type="NF",
        document_id=document_id,
        sequence="01",
        erp_entry_id=entry_id,
        invoice_number=invoice_number,
        document_value=value,
        other_discount=0,
        issue_date=purchase_date,
        due_date=due_date or payment_date or purchase_date,
        payment_date=payment_date,
        balance=0 if payment_date else value,
    ))
    if payment_date:
        db.add(PayableMovement(
            erp_key=f"B:{entry_id}",
            unit_code="777",
            person_id=77,
            title_type="NF",
            document_id=document_id,
            document_sequence="01",
            movement_type="B",
            movement_date=payment_date,
            amount=value,
            payment_sequence=1,
            reversal_sequence=0,
            financial_launch_id=entry_id,
            financial_sequence=1,
        ))
        db.add(FinancialEntry(
            erp_launch_id=entry_id,
            erp_sequence=1,
            unit_code="777",
            entry_date=payment_date,
            history_code=51,
            direction="S",
            value=value,
            document_id=document_id,
            history_text="PAGAMENTO IPIRANGA",
            origin="BPag",
        ))
    return purchase


def test_texaco_grouped_portal_credit_auto_confirms_exact_invariant_allocation():
    engine = create_engine("sqlite:///:memory:")
    Base.metadata.create_all(engine)
    with Session(engine) as db:
        rule = BonusRule(
            unit_code="777",
            company_code="TEXACO",
            kind="invoice_discount",
            effective_from=date(2025, 9, 1),
            rate_per_liter=Decimal("0.04"),
            applies_to="fuel_codes:1,3,5",
            active=True,
        )
        db.add(rule)
        first = _add_texaco_paid_purchase(
            db,
            entry_id=7701,
            invoice_number="770001",
            document_id="700001",
            liters="10000",
            purchase_date=date(2025, 9, 10),
            payment_date=date(2025, 9, 15),
        )
        _add_texaco_paid_purchase(
            db,
            entry_id=7702,
            invoice_number="770002",
            document_id="700002",
            liters="15000",
            purchase_date=date(2025, 9, 11),
            payment_date=date(2025, 9, 15),
        )
        event = PortalBonusEvent(
            id="texaco-group-1",
            unit_code="777",
            company_code="TEXACO",
            category="postpaid",
            portal_date=date(2025, 9, 16),
            value=Decimal("1000"),
            description="Nota Propria",
            event_key="texaco-group-key-1",
        )
        db.add(event)
        db.flush()

        assert match_ipiranga_events(db, "777", company_code="TEXACO") == 1
        match = db.scalar(select(PortalBonusMatch).where(PortalBonusMatch.event_id == event.id))
        details = json.loads(match.details_json)
        assert match.status == "portal_group_exact"
        assert {item["purchase_entry_id"] for item in details["allocations"]} == {7701, 7702}
        assert all(item["allocation_invariant"] for item in details["allocations"])

        evidence = _portal_invoice_credit_evidence(db, rule, first, Decimal("400"))
        assert evidence["portal_match_status"] == "portal_group_exact"
        assert evidence["source_value"] == 1000.0
        assert evidence["allocated"] == 400.0
        automatic, _ = _automatic_policy(
            rule,
            {
                "expected": Decimal("400"),
                "evidence": [evidence],
                "details": {"portal_postpaid_exact": True},
            },
            Decimal("400"),
        )
        assert automatic


def test_texaco_product_total_anchor_confirms_paid_title_and_keeps_unpaid_title_waiting():
    engine = create_engine("sqlite:///:memory:")
    Base.metadata.create_all(engine)
    with Session(engine) as db:
        rule = BonusRule(
            unit_code="777", company_code="TEXACO", kind="invoice_discount",
            effective_from=date(2026, 7, 1), rate_per_liter=Decimal("0.04"),
            applies_to="fuel_codes:1,3,5", active=True,
        )
        db.add(rule)
        paid = _add_texaco_paid_purchase(
            db, entry_id=7711, invoice_number="771101", document_id="711001", liters="10000",
            purchase_date=date(2026, 7, 8), payment_date=date(2026, 7, 9),
        )
        unpaid = _add_texaco_paid_purchase(
            db, entry_id=7712, invoice_number="771102", document_id="711002", liters="5000",
            purchase_date=date(2026, 7, 8), payment_date=None, due_date=date(2026, 7, 9),
        )
        imported = PortalStatementImport(
            unit_code="777", company_code="TEXACO", category="postpaid", client_cnpj="905896980012",
            period_start=date(2026, 7, 1), period_end=date(2026, 7, 9),
            original_filename="extrato.pdf", content_type="application/pdf", content_sha256="a" * 64,
            source_file=b"pdf", uploaded_by="admin",
        )
        event = PortalBonusEvent(
            id="texaco-anchor-1", unit_code="777", company_code="TEXACO", category="postpaid",
            portal_date=date(2026, 7, 9), value=Decimal("600"), description="Nota Propria",
            event_key="texaco-anchor-key-1",
        )
        db.add_all([imported, event])
        db.flush()
        db.add(PortalBonusEventSource(
            import_id=imported.id, event_id=event.id, row_number=1,
            raw_json='{"portal_product_total":"75000.00","portal_product_line_count":1}',
        ))
        db.flush()

        assert match_ipiranga_events(db, "777", company_code="TEXACO") == 1
        match = db.scalar(select(PortalBonusMatch).where(PortalBonusMatch.event_id == event.id))
        details = json.loads(match.details_json)
        assert match.status == "portal_issued_awaiting_payment"
        assert details["portal_product_total"] == "75000.00"
        assert {item["purchase_entry_id"] for item in details["allocations"]} == {7711, 7712}

        paid_evidence = _portal_invoice_credit_evidence(db, rule, paid, Decimal("400"))
        unpaid_evidence = _portal_invoice_credit_issued_evidence(db, rule, unpaid, Decimal("200"))
        assert paid_evidence["portal_product_anchor_exact"] is True
        assert paid_evidence["portal_settlement_confirmed"] is True
        assert unpaid_evidence["credit_issued"] is True
        assert unpaid_evidence["counted"] is True


def test_import_uses_active_texaco_rule_for_unit_007(monkeypatch):
    engine = create_engine("sqlite:///:memory:")
    Base.metadata.create_all(engine)
    parsed = type("Parsed", (), {
        "client_cnpj": "905896980006",
        "category_label": POSTPAID_CATEGORY,
        "period_start": date(2025, 7, 31),
        "period_end": date(2026, 1, 31),
        "events": [ParsedPortalEvent(1, date(2025, 9, 18), Decimal("600"), None, "Nota Propria", None, {})],
    })()
    monkeypatch.setattr("app.services.ipiranga_portal.parse_ipiranga_statement", lambda *_: parsed)
    with Session(engine) as db:
        db.add(BonusRule(
            unit_code="007",
            company_code="TEXACO",
            kind="invoice_discount",
            effective_from=date(2025, 9, 10),
            rate_per_liter=Decimal("0.04"),
            applies_to="fuel_codes:1,3,5",
            active=True,
        ))
        db.flush()
        imported, repeated = import_ipiranga_statement(
            db,
            unit_code="007",
            filename="extrato-007.pdf",
            content_type="application/pdf",
            content=b"unit-007",
            uploaded_by="admin",
        )
        assert not repeated
        assert imported.company_code == "TEXACO"


def test_import_allows_unit_014_only_with_the_validated_portal_cnpj_prefix(monkeypatch):
    engine = create_engine("sqlite:///:memory:")
    Base.metadata.create_all(engine)
    parsed = type("Parsed", (), {
        "client_cnpj": "00905896980012",
        "category_label": POSTPAID_CATEGORY,
        "period_start": date(2026, 4, 22),
        "period_end": date(2026, 7, 24),
        "events": [],
    })()
    monkeypatch.setattr("app.services.ipiranga_portal.parse_ipiranga_statement", lambda *_: parsed)
    with Session(engine) as db:
        db.add(BonusRule(
            unit_code="014", company_code="TEXACO", kind="invoice_discount",
            effective_from=date(2026, 4, 22), rate_per_liter=Decimal("0.04"),
            applies_to="fuel_codes:1,3,5", active=True,
        ))
        db.flush()
        imported, repeated = import_ipiranga_statement(
            db, unit_code="014", filename="extrato-014.pdf", content_type="application/pdf",
            content=b"unit-014", uploaded_by="admin",
        )
        assert not repeated
        assert imported.company_code == "TEXACO"


def test_cycle_ledger_reproduces_control_total_and_initial_shortfall():
    engine = create_engine("sqlite:///:memory:")
    Base.metadata.create_all(engine)
    cycle_liters = [115, 110, 115, 161, 135, 183, 149, 135, 159, 136, 143, 168, 135, 143, 140, 155, 174, 152]
    portal = [
        (date(2025, 5, 1), "14007"), (date(2025, 6, 3), "11745"),
        (date(2025, 7, 2), "15921"), (date(2025, 8, 1), "12963"),
        (date(2025, 9, 2), "22794"), (date(2025, 10, 2), "13833"),
        (date(2025, 11, 4), "11832"), (date(2025, 12, 9), "12441"),
        (date(2026, 1, 8), "14616.01"), (date(2026, 2, 2), "11745"),
        (date(2026, 3, 3), "12441"), (date(2026, 4, 2), "12180"),
        (date(2026, 5, 4), "13485"), (date(2026, 6, 1), "15138.01"),
        (date(2026, 7, 2), "13224"),
    ]
    with Session(engine) as db:
        seed_reference_data(db)
        rule = db.scalar(select(BonusRule).where(BonusRule.unit_code == "001"))
        start = rule.effective_from
        for index, liters in enumerate(cycle_liters, 1):
            issue_date = add_months(start, index - 1)
            db.add(
                Purchase(
                    erp_entry_id=10000 + index, unit_code="001", supplier_name="IPIRANGA",
                    supplier_cnpj="33337122015906", mapped_company_code="IPIRANGA",
                    invoice_number=str(index), purchase_date=issue_date, invoice_issue_date=issue_date,
                    total_liters=Decimal(liters * 1000), s10_liters=0, gross_value=1, net_value=1,
                )
            )
        for index, (portal_date, value) in enumerate(portal, 1):
            event = PortalBonusEvent(
                id=f"event-{index}", unit_code="001", company_code="IPIRANGA", category="postpaid",
                portal_date=portal_date, value=Decimal(value), description="Bonificação", event_key=f"key-{index}",
            )
            db.add(event)
            db.add(
                PortalBonusMatch(
                    event_id=event.id, status="exact", movement_key=f"D:{index}",
                    match_basis="exact", details_json='{"chosen": {}}', algorithm_version=ALGORITHM_VERSION,
                )
            )
        db.flush()
        ledger = build_ipiranga_cycle_ledger(db, rule, date(2026, 7, 13))
        expected = money(sum((row["expected"] for row in ledger["cycles"]), Decimal("0")))
        observed = money(sum((row["event"].value for row in ledger["events"]), Decimal("0")))
        assert expected == Decimal("226896.00")
        assert observed == Decimal("208365.02")
        assert expected - observed == Decimal("18530.98")
        assert ledger["cycles"][0]["observed"] == Decimal("10005.00")
        assert ledger["cycles"][1]["observed"] == Decimal("1044.02")
        assert ledger["cycles"][2]["observed"] == Decimal("0.00")
        assert all(row["difference"] == 0 for row in ledger["cycles"][3:])


def test_portal_value_exact_can_be_automatically_confirmed_with_auditable_evidence():
    engine = create_engine("sqlite:///:memory:")
    Base.metadata.create_all(engine)
    with Session(engine) as db:
        seed_reference_data(db)
        rule = db.scalar(select(BonusRule).where(BonusRule.unit_code == "001"))
        payload = {
            "expected": Decimal("14007"),
            "evidence": [{"source": "IPIRANGA_PORTAL", "value": 14007, "allocated": 14007, "chain_exact": True}],
        }
        automatic, _ = _automatic_policy(rule, payload, Decimal("14007"))
        assert automatic
        payload["evidence"][0]["chain_exact"] = False
        automatic, _ = _automatic_policy(rule, payload, Decimal("14007"))
        assert automatic
